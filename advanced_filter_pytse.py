#!/usr/bin/env python3
"""
فیلتر پیشرفته بورس بدون نیاز به سرور
ذخیره نتایج در Excel
"""

from pytse_client import Ticker
import pandas as pd
from datetime import datetime
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
import os


class AdvancedBourseFilter:
    """فیلتر پیشرفته بورس با قابلیت ذخیره در Excel"""

    def __init__(self, symbols_file='symbols.json', excel_file='filter_results.xlsx'):
        """مقداردهی اولیه"""
        self.symbols = self._load_symbols(symbols_file)
        self.excel_file = excel_file
        print(f"✅ {len(self.symbols)} نماد بارگذاری شد")

    def _load_symbols(self, symbols_file):
        """خواندن نمادها از فایل JSON"""
        try:
            with open(symbols_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                return [s['ticker'] for s in data.get('symbols', [])]
        except Exception as e:
            print(f"❌ خطا در خواندن فایل نمادها: {e}")
            return []

    def get_symbol_data(self, symbol):
        """دریافت داده‌های کامل یک نماد"""
        try:
            ticker = Ticker(symbol)

            # اطلاعات پایه
            data = {
                'نماد': symbol,
                'نام_شرکت': ticker.title if hasattr(ticker, 'title') else 'N/A',
                'قیمت_پایانی': ticker.adj_close,
                'قیمت_آخرین': ticker.last_price,
                'تغییر_درصد': 0,  # محاسبه می‌شود
                'حجم_معاملات': ticker.volume,
                'ارزش_معاملات': ticker.value,
                'تعداد_معاملات': ticker.count if hasattr(ticker, 'count') else 0,
            }

            # محاسبه تغییر درصد
            if hasattr(ticker, 'yesterday_price') and ticker.yesterday_price > 0:
                data['تغییر_درصد'] = round(
                    ((ticker.adj_close - ticker.yesterday_price) / ticker.yesterday_price) * 100,
                    2
                )

            # اطلاعات حقیقی/حقوقی
            client_types = ticker.client_types

            if client_types is not None and not client_types.empty:
                latest = client_types.iloc[-1]

                # حجم‌ها
                buy_I_Volume = latest.get('buy_I_Volume', 0)  # حقیقی خرید
                sell_I_Volume = latest.get('sell_I_Volume', 0)  # حقیقی فروش
                buy_N_Volume = latest.get('buy_N_Volume', 0)  # حقوقی خرید
                sell_N_Volume = latest.get('sell_N_Volume', 0)  # حقوقی فروش

                # محاسبات
                net_real_volume = buy_I_Volume - sell_I_Volume
                net_real_value = net_real_volume * ticker.last_price / 10_000_000

                total_real_volume = buy_I_Volume + sell_I_Volume
                real_power = (buy_I_Volume / total_real_volume * 100) if total_real_volume > 0 else 50.0

                data.update({
                    'خرید_حقیقی_حجم': buy_I_Volume,
                    'فروش_حقیقی_حجم': sell_I_Volume,
                    'خرید_حقوقی_حجم': buy_N_Volume,
                    'فروش_حقوقی_حجم': sell_N_Volume,
                    'ورود_پول_حقیقی_میلیون': round(net_real_value, 2),
                    'قدرت_خرید_حقیقی_درصد': round(real_power, 2),
                    'خالص_حجم_حقیقی': net_real_volume,
                })
            else:
                data.update({
                    'خرید_حقیقی_حجم': 0,
                    'فروش_حقیقی_حجم': 0,
                    'خرید_حقوقی_حجم': 0,
                    'فروش_حقوقی_حجم': 0,
                    'ورود_پول_حقیقی_میلیون': 0,
                    'قدرت_خرید_حقیقی_درصد': 0,
                    'خالص_حجم_حقیقی': 0,
                })

            return data

        except Exception as e:
            print(f"❌ خطا در دریافت اطلاعات {symbol}: {e}")
            return None

    def scan_all_symbols(self):
        """اسکن تمام نمادها"""
        print("\n🔍 در حال اسکن تمام نمادها...")

        results = []
        for i, symbol in enumerate(self.symbols, 1):
            print(f"[{i}/{len(self.symbols)}] {symbol}...")
            data = self.get_symbol_data(symbol)
            if data:
                results.append(data)

        return results

    def filter_data(self, data, filter_type, threshold):
        """
        اعمال فیلتر روی داده‌ها

        Args:
            data: لیست داده‌ها
            filter_type: نوع فیلتر ('money_flow', 'buying_power', 'price_change')
            threshold: آستانه فیلتر
        """
        if filter_type == 'money_flow':
            filtered = [d for d in data if d['ورود_پول_حقیقی_میلیون'] >= threshold]
            filtered.sort(key=lambda x: x['ورود_پول_حقیقی_میلیون'], reverse=True)

        elif filter_type == 'buying_power':
            filtered = [d for d in data if d['قدرت_خرید_حقیقی_درصد'] >= threshold]
            filtered.sort(key=lambda x: x['قدرت_خرید_حقیقی_درصد'], reverse=True)

        elif filter_type == 'price_change':
            filtered = [d for d in data if d['تغییر_درصد'] >= threshold]
            filtered.sort(key=lambda x: x['تغییر_درصد'], reverse=True)

        else:
            filtered = data

        return filtered

    def save_to_excel(self, data, sheet_name=None):
        """ذخیره داده‌ها در Excel"""
        if not data:
            print("❌ هیچ داده‌ای برای ذخیره وجود ندارد")
            return

        # ایجاد DataFrame
        df = pd.DataFrame(data)

        # نام شیت
        if sheet_name is None:
            sheet_name = datetime.now().strftime('%Y-%m-%d_%H-%M')

        # محدود کردن طول نام شیت به 31 کاراکتر
        if len(sheet_name) > 31:
            sheet_name = sheet_name[:31]

        try:
            # بارگذاری یا ایجاد فایل Excel
            if os.path.exists(self.excel_file):
                book = load_workbook(self.excel_file)
            else:
                book = Workbook()
                # حذف شیت پیش‌فرض
                if 'Sheet' in book.sheetnames:
                    del book['Sheet']

            # ایجاد شیت جدید
            ws = book.create_sheet(sheet_name)

            # اضافه کردن داده‌ها
            for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
                for c_idx, value in enumerate(row, 1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=value)

                    # فرمت‌بندی هدر
                    if r_idx == 1:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        # فرمت اعداد
                        if isinstance(value, (int, float)):
                            cell.number_format = '#,##0.00'

            # تنظیم عرض ستون‌ها
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 30)
                ws.column_dimensions[column_letter].width = adjusted_width

            # ذخیره فایل
            book.save(self.excel_file)
            print(f"✅ داده‌ها در شیت '{sheet_name}' ذخیره شد")
            print(f"📁 فایل: {self.excel_file}")

        except Exception as e:
            print(f"❌ خطا در ذخیره Excel: {e}")

    def print_summary(self, data):
        """چاپ خلاصه داده‌ها"""
        if not data:
            print("❌ داده‌ای موجود نیست")
            return

        print(f"\n{'='*100}")
        print(f"📊 خلاصه نتایج ({len(data)} نماد)")
        print(f"{'='*100}")

        # محاسبه آمار
        df = pd.DataFrame(data)

        print(f"\n💰 ورود پول حقیقی:")
        print(f"  • مجموع: {df['ورود_پول_حقیقی_میلیون'].sum():,.2f} میلیون تومان")
        print(f"  • میانگین: {df['ورود_پول_حقیقی_میلیون'].mean():,.2f} میلیون تومان")
        print(f"  • بیشترین: {df['ورود_پول_حقیقی_میلیون'].max():,.2f} میلیون تومان")

        print(f"\n📈 قدرت خرید حقیقی:")
        print(f"  • میانگین: {df['قدرت_خرید_حقیقی_درصد'].mean():.2f}%")
        print(f"  • بیشترین: {df['قدرت_خرید_حقیقی_درصد'].max():.2f}%")

        print(f"\n💹 تغییر قیمت:")
        print(f"  • میانگین: {df['تغییر_درصد'].mean():.2f}%")
        print(f"  • بیشترین: {df['تغییر_درصد'].max():.2f}%")
        print(f"  • کمترین: {df['تغییر_درصد'].min():.2f}%")

        # نمایش Top 5
        print(f"\n🏆 5 نماد برتر از نظر ورود پول:")
        top_5 = sorted(data, key=lambda x: x['ورود_پول_حقیقی_میلیون'], reverse=True)[:5]
        for i, item in enumerate(top_5, 1):
            print(f"  {i}. {item['نماد']:8s} → {item['ورود_پول_حقیقی_میلیون']:>10,.2f} م.تومان "
                  f"(قدرت: {item['قدرت_خرید_حقیقی_درصد']:.2f}%)")


def main():
    """تابع اصلی"""
    print("="*100)
    print("🚀 فیلتر پیشرفته بورس ایران - بدون نیاز به سرور + ذخیره Excel")
    print("="*100)

    # ایجاد فیلتر
    filter_tool = AdvancedBourseFilter()

    if not filter_tool.symbols:
        print("❌ هیچ نمادی یافت نشد")
        return

    # اسکن تمام نمادها
    all_data = filter_tool.scan_all_symbols()

    if not all_data:
        print("❌ هیچ داده‌ای دریافت نشد")
        return

    # نمایش منو
    while True:
        print("\n" + "="*100)
        print("📋 منوی اصلی:")
        print("="*100)
        print("1. نمایش تمام نمادها")
        print("2. فیلتر ورود پول (حقیقی)")
        print("3. فیلتر قدرت خرید (حقیقی)")
        print("4. فیلتر تغییر قیمت (درصد)")
        print("5. ذخیره نتایج فعلی در Excel")
        print("6. نمایش خلاصه آماری")
        print("7. اسکن مجدد")
        print("0. خروج")

        choice = input("\n➤ انتخاب شما: ").strip()

        if choice == "1":
            filter_tool.print_summary(all_data)
            print(f"\n✅ تمام {len(all_data)} نماد آماده است")

        elif choice == "2":
            threshold = input("حداقل ورود پول (میلیون تومان) [پیش‌فرض: 50]: ").strip()
            threshold = float(threshold) if threshold else 50.0

            filtered = filter_tool.filter_data(all_data, 'money_flow', threshold)
            filter_tool.print_summary(filtered)

            if filtered:
                save = input("\n💾 ذخیره در Excel? (y/n): ").strip().lower()
                if save == 'y':
                    sheet_name = f"ورود_پول_{threshold}"
                    filter_tool.save_to_excel(filtered, sheet_name)

        elif choice == "3":
            threshold = input("حداقل قدرت خرید (درصد) [پیش‌فرض: 55]: ").strip()
            threshold = float(threshold) if threshold else 55.0

            filtered = filter_tool.filter_data(all_data, 'buying_power', threshold)
            filter_tool.print_summary(filtered)

            if filtered:
                save = input("\n💾 ذخیره در Excel? (y/n): ").strip().lower()
                if save == 'y':
                    sheet_name = f"قدرت_خرید_{threshold}"
                    filter_tool.save_to_excel(filtered, sheet_name)

        elif choice == "4":
            threshold = input("حداقل تغییر قیمت (درصد) [پیش‌فرض: 2]: ").strip()
            threshold = float(threshold) if threshold else 2.0

            filtered = filter_tool.filter_data(all_data, 'price_change', threshold)
            filter_tool.print_summary(filtered)

            if filtered:
                save = input("\n💾 ذخیره در Excel? (y/n): ").strip().lower()
                if save == 'y':
                    sheet_name = f"تغییر_قیمت_{threshold}"
                    filter_tool.save_to_excel(filtered, sheet_name)

        elif choice == "5":
            sheet_name = input("نام شیت [Enter = تاریخ و زمان فعلی]: ").strip()
            if not sheet_name:
                sheet_name = None
            filter_tool.save_to_excel(all_data, sheet_name)

        elif choice == "6":
            filter_tool.print_summary(all_data)

        elif choice == "7":
            print("\n🔄 در حال اسکن مجدد...")
            all_data = filter_tool.scan_all_symbols()
            print("✅ اسکن مجدد کامل شد")

        elif choice == "0":
            print("\n👋 خداحافظ!")
            break

        else:
            print("❌ انتخاب نامعتبر!")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n\n⚠️  عملیات توسط کاربر لغو شد")
    except Exception as e:
        print(f"\n❌ خطای غیرمنتظره: {e}")
