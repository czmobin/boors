"""
مدیریت ذخیره داده در Excel - هر نماد یک شیت جداگانه
"""

import pandas as pd
import os
import json
from datetime import datetime
from typing import List, Dict, Set
from openpyxl import load_workbook, Workbook
from config import EXCEL_OUTPUT_DIR, EXCEL_FILENAME_TEMPLATE


class ExcelManager:
    """کلاس برای مدیریت ذخیره و بارگذاری داده‌ها در Excel"""

    def __init__(self, symbols_file: str = 'symbols.json'):
        # ایجاد پوشه خروجی در صورت عدم وجود
        if not os.path.exists(EXCEL_OUTPUT_DIR):
            os.makedirs(EXCEL_OUTPUT_DIR)

        self.current_date = datetime.now().strftime('%Y-%m-%d')
        self.filename = os.path.join(
            EXCEL_OUTPUT_DIR,
            EXCEL_FILENAME_TEMPLATE.format(date=self.current_date)
        )

        # بارگذاری لیست نمادهای مجاز از symbols.json
        self.allowed_symbols = self._load_allowed_symbols(symbols_file)

    def _load_allowed_symbols(self, symbols_file: str) -> Set[str]:
        """
        بارگذاری لیست نمادهای مجاز از فایل JSON
        فقط برای این نمادها شیت ساخته می‌شود
        """
        try:
            with open(symbols_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                symbols = data.get('symbols', [])

                # استخراج ticker نمادها
                allowed = set()
                for symbol in symbols:
                    ticker = symbol.get('ticker', '').strip()
                    if ticker:
                        allowed.add(ticker)

                print(f"📋 نمادهای مجاز برای Excel: {allowed}")
                return allowed

        except Exception as e:
            print(f"⚠️ خطا در بارگذاری symbols.json: {e}")
            return set()

    def save_data(self, data_list: List[Dict]) -> None:
        """
        ذخیره داده‌ها در Excel
        فقط برای نمادهای موجود در symbols.json شیت ساخته می‌شود
        هر نماد = یک شیت جداگانه
        هر اسکن = یک رکورد جدید در شیت مربوطه

        Args:
            data_list: لیستی از دیکشنری‌های حاوی داده‌های نمادها
        """
        if not data_list:
            print("هیچ داده‌ای برای ذخیره وجود ندارد")
            return

        if not self.allowed_symbols:
            print("⚠️ هیچ نمادی در symbols.json تعریف نشده - Excel ذخیره نمی‌شود")
            return

        try:
            current_time = datetime.now().strftime('%H:%M:%S')

            # بررسی وجود فایل
            if os.path.exists(self.filename):
                # فایل موجود است - append کن
                book = load_workbook(self.filename)
            else:
                # فایل جدید
                book = Workbook()
                # حذف شیت پیش‌فرض
                if 'Sheet' in book.sheetnames:
                    del book['Sheet']

            saved_count = 0

            # فقط برای نمادهای مجاز
            for item in data_list:
                symbol_name = item.get('نماد', 'N/A')

                # چک کن که این نماد در لیست مجازها هست
                if symbol_name not in self.allowed_symbols:
                    continue

                # ساخت رکورد بدون نام نماد (چون اسم شیت خودش نام نماده)
                record = self._prepare_record(item, current_time)

                # اگر شیت برای این نماد وجود داره
                if symbol_name in book.sheetnames:
                    # اضافه کردن رکورد جدید
                    self._append_record(book, symbol_name, record)
                else:
                    # ساخت شیت جدید
                    self._create_new_sheet(book, symbol_name, record)

                saved_count += 1

            # ذخیره فایل
            book.save(self.filename)
            print(f"✅ {saved_count} نماد در Excel ذخیره شد - زمان: {current_time}")

        except Exception as e:
            print(f"خطا در ذخیره داده‌ها: {e}")
            import traceback
            traceback.print_exc()

    def _prepare_record(self, item: Dict, current_time: str) -> Dict:
        """
        آماده‌سازی رکورد با فرمت صحیح اعداد (2 رقم اعشار)
        حذف فیلد 'نماد' از رکورد
        """
        record = {
            'زمان': current_time,
            'کد': str(item.get('کد', 'N/A')),
            'قیمت_پایانی': round(float(item.get('قیمت_پایانی', 0)), 2),
            'درصد_تغییر': round(float(item.get('درصد_تغییر', 0)), 2),
            'حجم_معاملات': int(item.get('حجم_معاملات', 0)),
            'ارزش_معاملات': int(item.get('ارزش_معاملات', 0)),
            'حجم_خرید_حقیقی': int(item.get('حجم_خرید_حقیقی', 0)),
            'تعداد_خرید_حقیقی': int(item.get('تعداد_خرید_حقیقی', 0)),
            'حجم_فروش_حقیقی': int(item.get('حجم_فروش_حقیقی', 0)),
            'تعداد_فروش_حقیقی': int(item.get('تعداد_فروش_حقیقی', 0)),
            'حجم_خرید_حقوقی': int(item.get('حجم_خرید_حقوقی', 0)),
            'حجم_فروش_حقوقی': int(item.get('حجم_فروش_حقوقی', 0)),
            'تعداد_خرید_حقوقی': int(item.get('تعداد_خرید_حقوقی', 0)),
            'تعداد_فروش_حقوقی': int(item.get('تعداد_فروش_حقوقی', 0)),
            'X_سرانه_خرید_حقیقی_تومان': round(float(item.get('X_سرانه_خرید_حقیقی_تومان', 0)), 2),
            'Y_سرانه_فروش_حقیقی_تومان': round(float(item.get('Y_سرانه_فروش_حقیقی_تومان', 0)), 2),
            'تفاضل_سرانه_تومان': round(float(item.get('تفاضل_سرانه_تومان', 0)), 2),
            'قدرت_خریدار': round(float(item.get('قدرت_خریدار', 0)), 2),
            'ورود_پول_حقوقی_میلیون': round(float(item.get('ورود_پول_حقوقی_میلیون', 0)), 2),
            'ورود_پول_حقیقی_میلیون': round(float(item.get('ورود_پول_حقیقی_میلیون', 0)), 2),
            'سرانه_خرید_حقوقی_میلیون': round(float(item.get('سرانه_خرید_حقوقی_میلیون', 0)), 2),
            'سرانه_فروش_حقوقی_میلیون': round(float(item.get('سرانه_فروش_حقوقی_میلیون', 0)), 2),
            'ورود_پول_خالص_میلیون': round(float(item.get('ورود_پول_خالص_میلیون', 0)), 2),
        }
        return record

    def _create_new_sheet(self, book: Workbook, sheet_name: str, record: Dict) -> None:
        """
        ساخت شیت جدید برای نماد با هدر و اولین رکورد
        """
        # محدود کردن طول نام شیت (Excel حداکثر 31 کاراکتر قبول میکنه)
        sheet_name = sheet_name[:31]

        sheet = book.create_sheet(title=sheet_name)

        # نوشتن هدر
        headers = list(record.keys())
        sheet.append(headers)

        # نوشتن اولین رکورد
        values = list(record.values())
        sheet.append(values)

    def _append_record(self, book: Workbook, sheet_name: str, record: Dict) -> None:
        """
        اضافه کردن رکورد جدید به شیت موجود
        """
        sheet = book[sheet_name]

        # نوشتن رکورد جدید
        values = list(record.values())
        sheet.append(values)

    def load_first_data(self) -> pd.DataFrame:
        """
        بارگذاری اولین داده‌های روز (برای مقایسه)
        از اولین ردیف هر شیت استفاده می‌کند

        Returns:
            DataFrame حاوی اولین داده‌های روز یا None
        """
        try:
            if not os.path.exists(self.filename):
                return None

            # خواندن همه شیت‌ها
            xls = pd.ExcelFile(self.filename)
            all_data = []

            for sheet_name in xls.sheet_names:
                if sheet_name == 'فیلتر_شده':
                    continue

                df = pd.read_excel(xls, sheet_name=sheet_name)

                if not df.empty:
                    # اولین ردیف = اولین رکورد روز
                    first_record = df.iloc[0].to_dict()
                    first_record['نماد'] = sheet_name  # اضافه کردن نام نماد
                    all_data.append(first_record)

            if all_data:
                return pd.DataFrame(all_data)

            return None

        except Exception as e:
            print(f"خطا در بارگذاری داده‌های اولیه: {e}")
            return None

    def create_summary_sheet(self, filtered_data: List[Dict]) -> None:
        """
        ایجاد شیت خلاصه برای نمادهای فیلتر شده

        Args:
            filtered_data: لیست نمادهای فیلتر شده
        """
        if not filtered_data:
            return

        try:
            # تبدیل به DataFrame
            df = pd.DataFrame(filtered_data)

            # فرمت کردن اعداد اعشاری
            numeric_columns = df.select_dtypes(include=['float64']).columns
            for col in numeric_columns:
                df[col] = df[col].round(2)

            # اگر فایل وجود داره، شیت فیلتر شده رو اضافه یا جایگزین کن
            if os.path.exists(self.filename):
                with pd.ExcelWriter(self.filename, mode='a', engine='openpyxl', if_sheet_exists='replace') as writer:
                    df.to_excel(writer, sheet_name='فیلتر_شده', index=False)
                print("✅ شیت 'فیلتر_شده' اضافه شد")
            else:
                print("⚠️ فایل Excel موجود نیست - ابتدا اسکن کنید")

        except Exception as e:
            print(f"خطا در ایجاد شیت خلاصه: {e}")
